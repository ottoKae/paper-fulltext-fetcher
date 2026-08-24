#!/usr/bin/env python3
"""Privacy-safe scholarly full-text acquisition helper.

Offline operations use the Python standard library. Excel input and robust PDF
text/page validation use optional openpyxl and pypdf dependencies. Network
commands are explicit, serial, rate-limited, do not use browser cookies, and
never print credential values.
"""

from __future__ import annotations

import argparse
import contextlib
import csv
import hashlib
import html
import ipaddress
import io
import json
import os
import re
import shutil
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

try:
    import tomllib
except ModuleNotFoundError:  # pragma: no cover - Python < 3.11 is rejected by doctor
    tomllib = None

try:
    import openpyxl
except ModuleNotFoundError:
    openpyxl = None

try:
    from pypdf import PdfReader
except ModuleNotFoundError:
    PdfReader = None


VERSION = "0.3.0"
STOP_HTTP = {403, 418, 429}
SUCCESS_STATUSES = {"valid_pdf", "valid_xml", "visual_approved_pdf", "already_present"}
APPROVED_FETCH_SCOPES = {"open_access", "institutional_repository", "official_api"}
WORKSPACE_GITIGNORE = """acquisition.local.toml
.env
papers_pdf/
papers_xml/
corpus/
manual_pdf_inbox/
outputs/
*.jsonl
*.har
*.sqlite*
*.xlsx
*.xlsm
bibliography*.csv
"""
TEXT_SUFFIXES = {
    ".csv", ".md", ".py", ".ps1", ".toml", ".yaml", ".yml", ".json",
    ".txt", ".html", ".gitignore", ".ini", ".cfg",
}
RISKY_SUFFIXES = {".pdf", ".xml", ".har", ".jsonl", ".sqlite", ".sqlite3", ".xlsx", ".xlsm", ".key", ".token", ".pem", ".p12", ".pfx"}
SKIP_DIRS = {".git", ".venv", "venv", "__pycache__", ".pytest_cache", "node_modules"}

ALIASES = {
    "title": {"title", "article title", "document title", "paper title"},
    "doi": {"doi", "input doi", "resolved doi", "digital object identifier"},
    "year": {"year", "publication year", "published", "publication date"},
    "source_title": {"source title", "journal", "publication name", "source"},
    "publisher": {"publisher", "publisher name"},
    "authors": {"authors", "author full names", "author"},
    "record_id": {"record id", "record_id", "uid", "ut unique wos id"},
    "expected_filename": {"expected filename", "expected_filename", "filename"},
    "excel_row": {"excel row", "excel_row", "row"},
}


def utcnow() -> str:
    return datetime.now(timezone.utc).isoformat()


def normalize_header(value: Any) -> str:
    return re.sub(r"[_\s]+", " ", str(value or "").strip().casefold())


def clean_doi(value: Any) -> str:
    text = str(value or "").strip().casefold()
    text = re.sub(r"^(?:https?://(?:dx\.)?doi\.org/|doi:\s*)", "", text)
    text = text.split("?", 1)[0].split("#", 1)[0].strip().strip(".;,()[]{}")
    match = re.search(r"10\.\d{4,9}/\S+", text)
    return match.group(0).rstrip(".;,)") if match else ""


def slug(value: str, limit: int = 72) -> str:
    text = re.sub(r"[^a-z0-9]+", "-", value.casefold()).strip("-")
    return (text[:limit].rstrip("-") or "untitled")


def identity_tokens(value: str) -> set[str]:
    tokens = re.findall(r"[a-z0-9]{3,}", value.casefold())
    stop = {"the", "and", "for", "with", "from", "using", "based", "study", "analysis"}
    return {token for token in tokens if token not in stop}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def stable_record_id(title: str, doi: str, year: str) -> str:
    basis = doi or f"{title.casefold().strip()}|{year.strip()}"
    return hashlib.sha256(basis.encode("utf-8")).hexdigest()[:12]


def expected_filename(record: dict[str, str], suffix: str = ".pdf") -> str:
    explicit = record.get("expected_filename", "").strip()
    if explicit:
        stem = Path(explicit).stem
    else:
        stem = f"{record['record_id']}_{slug(record['title'])}"
    return f"{stem}{suffix}"


def parse_column_overrides(values: Iterable[str]) -> dict[str, str]:
    overrides: dict[str, str] = {}
    for value in values:
        if "=" not in value:
            raise ValueError(f"invalid --column {value!r}; expected canonical=source_header")
        canonical, source = (part.strip() for part in value.split("=", 1))
        canonical = canonical.replace("-", "_")
        if canonical not in ALIASES:
            raise ValueError(f"unsupported canonical column {canonical!r}; choose from {sorted(ALIASES)}")
        if not source:
            raise ValueError(f"empty source header for {canonical}")
        overrides[canonical] = source
    return overrides


def canonical_value(row: dict[str, Any], field: str, column_overrides: dict[str, str] | None = None) -> str:
    aliases = ALIASES[field]
    normalized = {normalize_header(key): value for key, value in row.items()}
    if column_overrides and field in column_overrides:
        value = normalized.get(normalize_header(column_overrides[field]))
        return str(value).strip() if value is not None else ""
    for alias in aliases:
        value = normalized.get(normalize_header(alias))
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def canonical_record(
    row: dict[str, Any], default_row: int,
    column_overrides: dict[str, str] | None = None,
) -> dict[str, str]:
    title = canonical_value(row, "title", column_overrides)
    doi = clean_doi(canonical_value(row, "doi", column_overrides))
    year = canonical_value(row, "year", column_overrides)
    record_id = canonical_value(row, "record_id", column_overrides) or stable_record_id(title, doi, year)
    excel_row = canonical_value(row, "excel_row", column_overrides) or str(default_row)
    record = {
        "record_id": record_id,
        "excel_row": excel_row,
        "title": title,
        "doi": doi,
        "year": year,
        "source_title": canonical_value(row, "source_title", column_overrides),
        "publisher": canonical_value(row, "publisher", column_overrides) or "Unclassified",
        "authors": canonical_value(row, "authors", column_overrides),
        "expected_filename": canonical_value(row, "expected_filename", column_overrides),
    }
    record["expected_filename"] = expected_filename(record)
    return record


def read_table(path: Path, sheet_name: str | None = None) -> tuple[list[str], list[dict[str, Any]], dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(path)
    rows: list[dict[str, Any]] = []
    metadata: dict[str, Any] = {"input_type": path.suffix.casefold().lstrip("."), "sheet": "", "available_sheets": []}
    if path.suffix.casefold() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            headers = list(reader.fieldnames or [])
            rows = list(reader)
    elif path.suffix.casefold() in {".xlsx", ".xlsm"}:
        if openpyxl is None:
            raise RuntimeError("Excel input requires openpyxl; install requirements.txt")
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        metadata["available_sheets"] = list(workbook.sheetnames)
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                available = ", ".join(workbook.sheetnames)
                workbook.close()
                raise ValueError(f"sheet {sheet_name!r} not found; available: {available}")
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active
        metadata["sheet"] = sheet.title
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value or "") for value in next(iterator)]
        except StopIteration:
            workbook.close()
            raise ValueError("spreadsheet is empty") from None
        rows = [dict(zip(headers, values)) for values in iterator]
        workbook.close()
    else:
        raise ValueError("input must be .csv, .xlsx, or .xlsm")
    return headers, rows, metadata


def load_records(
    path: Path, sheet_name: str | None = None,
    column_overrides: dict[str, str] | None = None,
) -> list[dict[str, str]]:
    _, rows, _ = read_table(path, sheet_name)
    records = [canonical_record(row, index + 2, column_overrides) for index, row in enumerate(rows)]
    if not records:
        raise ValueError("bibliography contains no records")
    missing_identity = [record["excel_row"] for record in records if not record["title"] and not record["doi"]]
    if missing_identity:
        raise ValueError(f"missing both title and DOI in {len(missing_identity)} records; first rows: {missing_identity[:5]}")
    return records


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fields is None:
        fields = list(rows[0]) if rows else ["record_id"]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def append_jsonl(path: Path, event: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(event, ensure_ascii=False) + "\n")


def redact_url(value: str) -> str:
    try:
        parsed = urllib.parse.urlsplit(value)
    except ValueError:
        return ""
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        return ""
    return urllib.parse.urlunsplit((parsed.scheme, parsed.netloc, parsed.path, "", ""))


def redact_error_message(value: str) -> str:
    """Remove developer/user identifiers from terminal error messages."""
    text = value
    for local in {str(Path.home()), str(Path.cwd())}:
        if local:
            text = text.replace(local, "<local-path>")
            text = text.replace(local.replace("\\", "/"), "<local-path>")
    text = re.sub(r"(?i)\b[A-Z]:[\\/](?:[^\\/\s'\"]+[\\/])*[^\\/\s'\"]*", "<local-path>", text)
    text = re.sub(r"/(?:home|Users)/[^/\s'\"]+(?:/[^\s'\"]*)?", "<local-path>", text)
    text = re.sub(r"(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", "<email>", text)
    text = re.sub(r"(?i)([?&](?:token|signature|sig|key|auth|expires)=)[^&\s]+", r"\1<redacted>", text)
    return text


def read_prefix(path: Path, limit: int = 1024) -> bytes:
    with path.open("rb") as handle:
        return handle.read(limit)


def pdf_text(path: Path, max_pages: int = 5) -> tuple[str, str]:
    if PdfReader is None:
        return "", "pypdf_not_installed"
    try:
        reader = PdfReader(path)
        parts = []
        for page in reader.pages[:max_pages]:
            parts.append(page.extract_text() or "")
        return "\n".join(parts), f"pages={len(reader.pages)}"
    except Exception as exc:  # pypdf raises several format-specific exceptions
        return "", f"pdf_parse_failed:{type(exc).__name__}"


def title_score(text: str, title: str) -> tuple[float, int]:
    wanted = identity_tokens(title)
    observed = identity_tokens(text)
    matched = wanted & observed
    return (len(matched) / len(wanted) if wanted else 0.0), len(matched)


def validate_pdf(path: Path, record: dict[str, str], minimum_bytes: int) -> tuple[bool, str, str]:
    if not path.exists():
        return False, "missing", ""
    if path.stat().st_size < minimum_bytes:
        return False, "below_minimum_bytes", ""
    if not read_prefix(path).lstrip().startswith(b"%PDF-"):
        return False, "invalid_pdf_signature", ""
    text, parse_detail = pdf_text(path)
    lowered = text.casefold()
    doi = record["doi"]
    if doi and doi in lowered:
        return True, f"doi_match;{parse_detail}", text
    coverage, matches = title_score(text, record["title"]) if text else (0.0, 0)
    if coverage >= 0.70 and matches >= 5:
        return True, f"title_match={coverage:.2f};{parse_detail}", text
    return False, f"identity_mismatch:title={coverage:.2f};{parse_detail}", text


def validate_xml(path: Path, record: dict[str, str], minimum_bytes: int) -> tuple[bool, str]:
    if not path.exists() or path.stat().st_size < minimum_bytes:
        return False, "missing_or_too_small"
    text = path.read_text(encoding="utf-8", errors="ignore")[:2_000_000]
    lowered = text.casefold()
    if "<html" in lowered[:1000]:
        return False, "html_not_xml_fulltext"
    fulltext_markers = ("<body", "<ce:sections", "<article", "<sec", "<xocs:doc")
    if not any(marker in lowered for marker in fulltext_markers):
        return False, "no_fulltext_structure"
    if record["doi"] and record["doi"] in lowered:
        return True, "doi_match"
    coverage, matches = title_score(text, record["title"])
    if coverage >= 0.70 and matches >= 5:
        return True, f"title_match={coverage:.2f}"
    return False, f"identity_mismatch:title={coverage:.2f}"


def load_visual_approvals(paths: Iterable[Path]) -> dict[str, dict[str, Any]]:
    approvals: dict[str, dict[str, Any]] = {}
    for path in paths:
        if not path.exists():
            continue
        with path.open("r", encoding="utf-8-sig") as handle:
            for line in handle:
                try:
                    event = json.loads(line)
                except json.JSONDecodeError:
                    continue
                if event.get("status") not in {"visual_approved_imported", "visual_approved_pdf"}:
                    continue
                digest = str(event.get("sha256", "")).casefold()
                filename = str(event.get("expected_filename") or event.get("filename") or "").casefold()
                if digest and filename:
                    approvals[f"{filename}|{digest}"] = event
    return approvals


def match_file_by_text(path: Path, records: list[dict[str, str]], suffix: str, minimum_bytes: int) -> tuple[dict[str, str] | None, str]:
    if suffix == ".pdf":
        if path.stat().st_size < minimum_bytes or not read_prefix(path).lstrip().startswith(b"%PDF-"):
            return None, "invalid_container"
        text, detail = pdf_text(path)
    else:
        text = path.read_text(encoding="utf-8", errors="ignore")[:2_000_000]
        detail = "xml_text"
    lowered = text.casefold()
    doi_matches = [record for record in records if record["doi"] and record["doi"] in lowered]
    if len(doi_matches) == 1:
        return doi_matches[0], f"doi_match;{detail}"
    ranked = []
    for record in records:
        coverage, matches = title_score(text, record["title"])
        ranked.append((coverage, matches, record))
    ranked.sort(key=lambda item: (item[0], item[1]), reverse=True)
    if ranked and ranked[0][0] >= 0.70 and ranked[0][1] >= 5:
        if len(ranked) == 1 or ranked[0][0] - ranked[1][0] >= 0.05:
            return ranked[0][2], f"title_match={ranked[0][0]:.2f};{detail}"
        return None, "ambiguous_title_match"
    return None, "no_identity_match"


def audit_records(
    records: list[dict[str, str]], pdf_dir: Path, xml_dir: Path,
    minimum_bytes: int, approvals: dict[str, dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    matched: dict[str, dict[str, Any]] = {}
    anomalies: list[dict[str, Any]] = []
    all_files: list[Path] = []
    expected_index = {record["expected_filename"].casefold(): record for record in records}
    stem_index = {Path(record["expected_filename"]).stem.casefold(): record for record in records}

    for directory, suffix in ((pdf_dir, ".pdf"), (xml_dir, ".xml")):
        if not directory.exists():
            continue
        for path in sorted(directory.glob(f"*{suffix}")):
            all_files.append(path)
            record = expected_index.get(path.name.casefold()) or stem_index.get(path.stem.casefold())
            detail = "expected_filename"
            if record is None:
                record, detail = match_file_by_text(path, records, suffix, minimum_bytes)
            if record is None:
                anomalies.append({"path": path.name, "kind": "orphan_or_invalid", "detail": detail})
                continue
            if suffix == ".pdf":
                container_ok = path.stat().st_size >= minimum_bytes and read_prefix(path).lstrip().startswith(b"%PDF-")
                digest = sha256_file(path) if container_ok else ""
                approval_key = f"{path.name.casefold()}|{digest.casefold()}"
                visually_approved = container_ok and approval_key in approvals
                if visually_approved:
                    valid, validation, status = True, "visual_approval_hash_match", "visual_approved_pdf"
                else:
                    valid, validation, _ = validate_pdf(path, record, minimum_bytes)
                    status = "valid_pdf" if valid else "invalid_container"
            else:
                valid, validation = validate_xml(path, record, minimum_bytes)
                digest = sha256_file(path) if valid else ""
                status = "valid_xml" if valid else "invalid_container"
            if not valid:
                anomalies.append({"path": path.name, "kind": status, "detail": validation})
                continue
            existing = matched.get(record["record_id"])
            if existing:
                anomalies.append({"path": path.name, "kind": "multiple_files_for_record", "detail": existing["fulltext_file"]})
                if existing["status"] in {"valid_pdf", "visual_approved_pdf"}:
                    continue
            matched[record["record_id"]] = {
                "status": status, "fulltext_file": path.name, "validation": validation,
                "bytes": path.stat().st_size, "sha256": digest,
            }

    digest_groups: dict[str, list[str]] = defaultdict(list)
    for item in matched.values():
        digest_groups[item["sha256"]].append(item["fulltext_file"])
    for digest, paths in digest_groups.items():
        if digest and len(paths) > 1:
            anomalies.append({"path": " | ".join(paths), "kind": "duplicate_content", "detail": digest})

    output_rows: list[dict[str, Any]] = []
    for record in records:
        item = matched.get(record["record_id"], {})
        output_rows.append({
            **record,
            "status": item.get("status", "missing"),
            "fulltext_file": item.get("fulltext_file", ""),
            "validation": item.get("validation", ""),
            "bytes": item.get("bytes", ""),
            "sha256": item.get("sha256", ""),
        })
    counts = Counter(row["status"] for row in output_rows)
    coverage = sum(counts[status] for status in ("valid_pdf", "visual_approved_pdf", "valid_xml"))
    summary = {
        "finished_utc": utcnow(),
        "helper_version": VERSION,
        "bibliography_records": len(records),
        "valid_pdf_records": counts["valid_pdf"] + counts["visual_approved_pdf"],
        "visual_approved_pdf_records": counts["visual_approved_pdf"],
        "valid_xml_only_records": counts["valid_xml"],
        "fulltext_coverage": coverage,
        "missing_fulltext": counts["missing"],
        "coverage_percent": round(100 * coverage / len(records), 2),
        "files_scanned": len(all_files),
        "anomalies": len(anomalies),
        "status": dict(sorted(counts.items())),
    }
    return output_rows, anomalies, summary


def skill_root() -> Path:
    return Path(__file__).resolve().parents[1]


def resolve_local_path(value: str | Path, base: Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else base / path


def load_runtime_config(path: Path | None) -> tuple[dict[str, Any], Path]:
    if path is None:
        return {}, Path.cwd()
    config = load_config(path)
    return config, path.resolve().parent


def configured_path(
    args: argparse.Namespace, argument: str, config_key: str,
    *, default_relative: str | None = None,
) -> Path:
    explicit = getattr(args, argument, None)
    if explicit is not None:
        return Path(explicit)
    config, base = load_runtime_config(getattr(args, "config", None))
    project = config.get("project") or {}
    value = project.get(config_key, default_relative)
    if not value:
        raise ValueError(f"provide --{argument.replace('_', '-')} or set [project].{config_key} in --config")
    return resolve_local_path(str(value), base)


def configured_output_child(args: argparse.Namespace, argument: str, child: str) -> Path:
    explicit = getattr(args, argument, None)
    if explicit is not None:
        return Path(explicit)
    config, base = load_runtime_config(getattr(args, "config", None))
    output_root = str((config.get("project") or {}).get("output_root", "outputs/acquisition"))
    return resolve_local_path(output_root, base) / child


def runtime_columns(args: argparse.Namespace) -> dict[str, str]:
    config, _ = load_runtime_config(getattr(args, "config", None))
    configured = {str(key).replace("-", "_"): str(value) for key, value in (config.get("columns") or {}).items() if str(value).strip()}
    configured.update(parse_column_overrides(getattr(args, "column", []) or []))
    return configured


def runtime_sheet(args: argparse.Namespace) -> str | None:
    if getattr(args, "sheet", None):
        return str(args.sheet)
    config, _ = load_runtime_config(getattr(args, "config", None))
    value = (config.get("project") or {}).get("sheet")
    return str(value) if value else None


def command_init(args: argparse.Namespace) -> int:
    workspace = args.workspace.resolve()
    config_path = workspace / "acquisition.local.toml"
    if args.input.is_absolute():
        raise ValueError("--input must be relative so the generated config does not expose a developer path")
    if config_path.exists() and not args.force:
        print("acquisition.local.toml already exists; use --force only after reviewing it", file=sys.stderr)
        return 2
    template = (skill_root() / "assets" / "acquisition.example.toml").read_text(encoding="utf-8")
    input_value = json.dumps(args.input.as_posix(), ensure_ascii=False)
    template = re.sub(r'(?m)^input\s*=\s*"[^"]*"', f"input = {input_value}", template, count=1)
    if args.dry_run:
        print(json.dumps({
            "mode": "dry_run",
            "would_create": ["acquisition.local.toml", ".gitignore if absent", "papers_pdf", "papers_xml", "manual_pdf_inbox", "outputs/acquisition"],
        }, ensure_ascii=False, indent=2))
        return 0
    workspace.mkdir(parents=True, exist_ok=True)
    for relative in ("papers_pdf", "papers_xml", "manual_pdf_inbox", "outputs/acquisition"):
        (workspace / relative).mkdir(parents=True, exist_ok=True)
    config_path.write_text(template, encoding="utf-8")
    gitignore = workspace / ".gitignore"
    gitignore_status = "existing_not_modified"
    if not gitignore.exists():
        gitignore.write_text(WORKSPACE_GITIGNORE, encoding="utf-8")
        gitignore_status = "created"
    print(json.dumps({
        "status": "initialized",
        "config": config_path.name, "input_expected": args.input.as_posix(),
        "gitignore": gitignore_status,
        "next": ["copy the bibliography into the workspace", "run doctor", "run inspect-input", "run audit"],
    }, ensure_ascii=False, indent=2))
    return 0


def detected_columns(headers: list[str], overrides: dict[str, str]) -> dict[str, str]:
    normalized = {normalize_header(header): header for header in headers}
    mapping: dict[str, str] = {}
    for field, aliases in ALIASES.items():
        if field in overrides and normalize_header(overrides[field]) in normalized:
            mapping[field] = normalized[normalize_header(overrides[field])]
            continue
        for alias in aliases:
            if normalize_header(alias) in normalized:
                mapping[field] = normalized[normalize_header(alias)]
                break
    return mapping


def command_inspect_input(args: argparse.Namespace) -> int:
    input_path = configured_path(args, "input", "input")
    sheet = runtime_sheet(args)
    overrides = runtime_columns(args)
    headers, rows, metadata = read_table(input_path, sheet)
    records = [canonical_record(row, index + 2, overrides) for index, row in enumerate(rows)]
    mapping = detected_columns(headers, overrides)
    dois = [record["doi"] for record in records if record["doi"]]
    result = {
        "checked_at_utc": utcnow(),
        "input_type": metadata["input_type"],
        "sheet": metadata["sheet"],
        "available_sheets": metadata["available_sheets"],
        "records": len(records),
        "headers": len(headers),
        "mapped_columns": mapping,
        "unmapped_optional_fields": [field for field in ALIASES if field not in mapping and field != "title"],
        "records_with_title": sum(bool(record["title"]) for record in records),
        "records_with_doi": len(dois),
        "records_with_title_or_doi": sum(bool(record["title"] or record["doi"]) for record in records),
        "duplicate_doi_values": len(dois) - len(set(dois)),
        "unclassified_publishers": sum(record["publisher"] == "Unclassified" for record in records),
        "ready_for_audit": bool(records) and all(record["title"] or record["doi"] for record in records),
    }
    if args.output:
        write_json(args.output, result)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["ready_for_audit"] else 2


def command_audit(args: argparse.Namespace) -> int:
    args.input = configured_path(args, "input", "input")
    args.pdf_dir = configured_path(args, "pdf_dir", "pdf_dir")
    args.xml_dir = configured_path(args, "xml_dir", "xml_dir")
    args.output_dir = configured_path(args, "output_dir", "output_root", default_relative="outputs/acquisition") / "audit" if args.output_dir is None else args.output_dir
    records = load_records(args.input, runtime_sheet(args), runtime_columns(args))
    approvals = load_visual_approvals(args.visual_approvals)
    rows, anomalies, summary = audit_records(records, args.pdf_dir, args.xml_dir, args.minimum_pdf_bytes, approvals)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(args.output_dir / "records.csv", rows)
    write_csv(args.output_dir / "missing.csv", [row for row in rows if row["status"] == "missing"])
    write_csv(args.output_dir / "anomalies.csv", anomalies, ["path", "kind", "detail"])
    publisher_counts: dict[str, Counter] = defaultdict(Counter)
    for row in rows:
        publisher_counts[row["publisher"]][row["status"]] += 1
    publisher_rows = []
    for publisher, counts in sorted(publisher_counts.items()):
        total = sum(counts.values())
        covered = total - counts["missing"]
        publisher_rows.append({
            "publisher": publisher, "total": total,
            "valid_pdf": counts["valid_pdf"] + counts["visual_approved_pdf"],
            "valid_xml": counts["valid_xml"], "missing": counts["missing"],
            "coverage_percent": f"{100 * covered / total:.2f}",
        })
    write_csv(args.output_dir / "publisher_coverage.csv", publisher_rows)
    write_json(args.output_dir / "summary.json", summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if not anomalies else 2


def read_remaining(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    return [row for row in rows if not row.get("status") or row.get("status") == "missing"]


def route_for(publisher: str) -> str:
    name = publisher.casefold()
    if "elsevier" in name:
        return "OA; authorized Article Retrieval/DaaS; institutional manual"
    if "wiley" in name:
        return "OA; authorized Wiley TDM; institutional manual"
    if "taylor" in name or "francis" in name:
        return "OA; institution-approved TDM; institutional manual"
    if "ieee" in name:
        return "OA/API; official journal page; institutional manual"
    return "OA or repository; library delivery; author; institutional manual"


def command_manual_queue(args: argparse.Namespace) -> int:
    args.remaining = configured_output_child(args, "remaining", "audit/missing.csv")
    args.output_dir = configured_output_child(args, "output_dir", "manual")
    rows = read_remaining(args.remaining)
    output_rows = []
    for row in rows:
        doi = clean_doi(row.get("doi") or row.get("input_doi") or row.get("resolved_doi"))
        output_rows.append({
            **row,
            "doi": doi,
            "doi_landing": f"https://doi.org/{doi}" if doi else "",
            "recommended_route": route_for(row.get("publisher", "")),
            "manual_status": row.get("manual_status", "pending") or "pending",
        })
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(args.output_dir / "manual_queue.csv", output_rows)
    body = []
    for row in output_rows:
        link = html.escape(row.get("doi_landing", ""), quote=True)
        anchor = f'<a href="{link}" target="_blank" rel="noreferrer">open DOI</a>' if link else "no DOI"
        body.append(
            "<tr>" + "".join([
                f"<td>{html.escape(row.get('excel_row', ''))}</td>",
                f"<td>{html.escape(row.get('publisher', ''))}</td>",
                f"<td>{html.escape(row.get('year', ''))}</td>",
                f"<td>{html.escape(row.get('title', ''))}</td>",
                f"<td>{anchor}</td>",
                f"<td>{html.escape(row.get('expected_filename', ''))}</td>",
                f"<td>{html.escape(row.get('recommended_route', ''))}</td>",
            ]) + "</tr>"
        )
    page = (
        "<!doctype html><meta charset='utf-8'><title>Manual full-text queue</title>"
        "<style>body{font-family:system-ui;margin:24px}input{width:45%;padding:8px}"
        "table{border-collapse:collapse;width:100%;margin-top:12px}td,th{border:1px solid #ccc;padding:6px;vertical-align:top}"
        "th{position:sticky;top:0;background:#fff}tr:nth-child(even){background:#f7f7f7}</style>"
        "<h1>Manual full-text queue</h1><p>Use only lawful institutional, library, repository or author-provided access.</p>"
        "<input id='q' placeholder='Filter' oninput=\"for(const r of document.querySelectorAll('tbody tr'))"
        "r.hidden=!r.innerText.toLowerCase().includes(this.value.toLowerCase())\">"
        "<table><thead><tr><th>Row</th><th>Publisher</th><th>Year</th><th>Title</th><th>DOI</th>"
        "<th>Expected file</th><th>Route</th></tr></thead><tbody>" + "".join(body) + "</tbody></table>"
    )
    (args.output_dir / "manual_queue.html").write_text(page, encoding="utf-8")
    summary = {"created_utc": utcnow(), "records": len(output_rows), "publishers": dict(Counter(row.get("publisher", "") for row in output_rows))}
    write_json(args.output_dir / "summary.json", summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


def best_record_for_pdf(path: Path, records: list[dict[str, str]], minimum_bytes: int) -> tuple[dict[str, str] | None, str]:
    exact = [record for record in records if path.name.casefold() == record["expected_filename"].casefold()]
    if len(exact) == 1:
        valid, detail, _ = validate_pdf(path, exact[0], minimum_bytes)
        return (exact[0], detail) if valid else (None, detail)
    record, detail = match_file_by_text(path, records, ".pdf", minimum_bytes)
    if record is None:
        return None, detail
    valid, validation, _ = validate_pdf(path, record, minimum_bytes)
    return (record, validation) if valid else (None, validation)


def command_import(args: argparse.Namespace) -> int:
    args.input = configured_path(args, "input", "input")
    args.remaining = configured_output_child(args, "remaining", "audit/missing.csv")
    args.inbox = configured_path(args, "inbox", "manual_inbox")
    args.pdf_dir = configured_path(args, "pdf_dir", "pdf_dir")
    args.log_dir = configured_output_child(args, "log_dir", "import")
    all_records = load_records(args.input, runtime_sheet(args), runtime_columns(args))
    remaining_rows = read_remaining(args.remaining)
    wanted_ids = {row.get("record_id", "") for row in remaining_rows if row.get("record_id")}
    wanted_excel_rows = {str(row.get("excel_row", "")) for row in remaining_rows if row.get("excel_row")}
    records = [
        record for record in all_records
        if (not wanted_ids and not wanted_excel_rows)
        or record["record_id"] in wanted_ids
        or record["excel_row"] in wanted_excel_rows
    ]
    args.log_dir.mkdir(parents=True, exist_ok=True)
    events_path = args.log_dir / f"import_{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}.jsonl"
    results = []
    for source in sorted(args.inbox.glob("*.pdf")):
        record, detail = best_record_for_pdf(source, records, args.minimum_pdf_bytes)
        event: dict[str, Any] = {
            "timestamp_utc": utcnow(), "action": "manual_import", "source_file": source.name,
            "status": "", "validation": detail,
        }
        if record is None:
            event["status"] = "visual_approval_required" if "pypdf_not_installed" not in detail and read_prefix(source).lstrip().startswith(b"%PDF-") else "identity_mismatch"
        else:
            destination = args.pdf_dir / record["expected_filename"]
            event.update(record_id=record["record_id"], excel_row=record["excel_row"], title=record["title"], destination=destination.name)
            if destination.exists():
                event["status"] = "already_present"
            elif args.dry_run:
                event["status"] = "would_import"
            else:
                args.pdf_dir.mkdir(parents=True, exist_ok=True)
                temporary = destination.with_suffix(".pdf.part")
                shutil.copy2(source, temporary)
                temporary.replace(destination)
                event.update(status="valid_pdf", bytes=destination.stat().st_size, sha256=sha256_file(destination))
        append_jsonl(events_path, event)
        results.append(event)
        print(f"{source.name}: {event['status']}")
    summary = {"finished_utc": utcnow(), "dry_run": args.dry_run, "files": len(results), "status": dict(Counter(event["status"] for event in results)), "events": events_path.name}
    write_json(args.log_dir / "summary.json", summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if all(event["status"] in {"already_present", "would_import", "valid_pdf"} for event in results) else 2


def http_json(url: str, headers: dict[str, str], timeout: int) -> tuple[int, Any]:
    request = urllib.request.Request(url, headers={"User-Agent": "paper-fulltext-fetcher/0.1", **headers})
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return response.status, json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        return exc.code, None
    except (urllib.error.URLError, TimeoutError, json.JSONDecodeError):
        return 0, None


def oa_candidates_for(record: dict[str, str], email: str, openalex_key: str, timeout: int, delay: float, event_path: Path) -> tuple[list[dict[str, str]], bool]:
    doi = record["doi"]
    if not doi:
        return [], False
    candidates: list[dict[str, str]] = []
    stop = False
    services = []
    if email:
        services.append(("unpaywall", f"https://api.unpaywall.org/v2/{urllib.parse.quote(doi, safe='')}?email={urllib.parse.quote(email)}"))
    openalex_query = urllib.parse.urlencode({"api_key": openalex_key}) if openalex_key else ""
    openalex_url = f"https://api.openalex.org/works/{urllib.parse.quote('https://doi.org/' + doi, safe=':/')}"
    if openalex_query:
        openalex_url += "?" + openalex_query
    services.append(("openalex", openalex_url))
    services.append(("crossref", f"https://api.crossref.org/works/{urllib.parse.quote(doi, safe='')}"))

    for provider, url in services:
        status, payload = http_json(url, {}, timeout)
        event = {
            "timestamp_utc": utcnow(), "action": "discover_oa", "provider": provider,
            "record_id": record["record_id"], "doi": doi, "http_status": status,
            "request_url_redacted": redact_url(url),
            "status": "ok" if status == 200 else ("rate_limited" if status == 429 else "failed"),
        }
        append_jsonl(event_path, event)
        if status in STOP_HTTP:
            stop = True
            break
        if status == 200 and isinstance(payload, dict):
            locations: list[tuple[str, str, str]] = []
            if provider == "unpaywall":
                best = payload.get("best_oa_location") or {}
                locations.append((best.get("url_for_pdf") or "", best.get("url") or "", best.get("license") or ""))
            elif provider == "openalex":
                best = payload.get("best_oa_location") or payload.get("primary_location") or {}
                locations.append((best.get("pdf_url") or "", best.get("landing_page_url") or "", str((payload.get("open_access") or {}).get("oa_status") or "")))
            else:
                message = payload.get("message") or {}
                for link in message.get("link") or []:
                    locations.append((link.get("URL") or "", "", link.get("content-type") or ""))
            for pdf_url, landing, license_value in locations:
                chosen = pdf_url or landing
                safe_url = redact_url(chosen)
                if not safe_url:
                    continue
                candidate_record = dict(record)
                if "xml" in license_value.casefold():
                    candidate_record["expected_filename"] = expected_filename(record, ".xml")
                candidates.append({
                    **candidate_record, "provider": provider, "candidate_url": safe_url,
                    "landing_url": redact_url(landing), "license_or_type": license_value,
                    "approved": "false", "approval_scope": "open_access_candidate",
                })
        time.sleep(max(0.0, delay))
    return candidates, stop


def command_discover_oa(args: argparse.Namespace) -> int:
    args.remaining = configured_output_child(args, "remaining", "audit/missing.csv")
    args.output = configured_output_child(args, "output", "oa_candidates.csv")
    if args.delay < 1.0:
        print("--delay must be at least 1 second", file=sys.stderr)
        return 2
    rows = read_remaining(args.remaining)
    if args.limit:
        rows = rows[: args.limit]
    records = [canonical_record(row, int(row.get("excel_row") or index + 2)) for index, row in enumerate(rows)]
    email = os.getenv(args.unpaywall_email_env, "")
    openalex_key = os.getenv(args.openalex_key_env, "")
    if not email and not openalex_key:
        print("No OA credentials configured. Set UNPAYWALL_EMAIL and/or OPENALEX_API_KEY.", file=sys.stderr)
        return 2
    args.output.parent.mkdir(parents=True, exist_ok=True)
    event_path = args.output.with_suffix(".events.jsonl")
    candidates: list[dict[str, str]] = []
    stopped = False
    for record in records:
        found, stopped = oa_candidates_for(record, email, openalex_key, args.timeout, args.delay, event_path)
        candidates.extend(found)
        if stopped:
            break
    write_csv(args.output, candidates)
    summary = {"finished_utc": utcnow(), "records_requested": len(records), "candidates": len(candidates), "circuit_breaker_triggered": stopped, "events": event_path.name}
    write_json(args.output.with_suffix(".summary.json"), summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 3 if stopped else 0


def approved_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    return [
        row for row in rows
        if str(row.get("approved", "")).casefold() in {"true", "yes", "1"}
        and str(row.get("approval_scope", "")).casefold() in APPROVED_FETCH_SCOPES
    ]


def is_allowed_host(host: str, allowed_hosts: set[str]) -> bool:
    normalized = host.casefold().strip(".")
    if not normalized:
        return False
    try:
        address = ipaddress.ip_address(normalized)
    except ValueError:
        address = None
    if address and (address.is_private or address.is_loopback or address.is_link_local or address.is_reserved):
        return False
    return any(normalized == allowed or normalized.endswith("." + allowed) for allowed in allowed_hosts)


class RestrictedRedirectHandler(urllib.request.HTTPRedirectHandler):
    def __init__(self, allowed_hosts: set[str]) -> None:
        super().__init__()
        self.allowed_hosts = allowed_hosts

    def redirect_request(self, req: Any, fp: Any, code: int, msg: str, headers: Any, newurl: str) -> Any:
        host = urllib.parse.urlsplit(newurl).hostname or ""
        if not is_allowed_host(host, self.allowed_hosts):
            raise urllib.error.URLError("redirect_host_not_allowed")
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def download_one(url: str, temporary: Path, timeout: int, max_bytes: int, allowed_hosts: set[str]) -> tuple[int, str]:
    request = urllib.request.Request(url, headers={"User-Agent": "paper-fulltext-fetcher/0.1"})
    opener = urllib.request.build_opener(RestrictedRedirectHandler(allowed_hosts))
    try:
        with opener.open(request, timeout=timeout) as response, temporary.open("wb") as handle:
            total = 0
            while True:
                chunk = response.read(1024 * 1024)
                if not chunk:
                    break
                total += len(chunk)
                if total > max_bytes:
                    return response.status, "maximum_bytes_exceeded"
                handle.write(chunk)
            return response.status, str(response.headers.get("Content-Type", ""))
    except urllib.error.HTTPError as exc:
        return exc.code, "http_error"
    except (urllib.error.URLError, TimeoutError):
        return 0, "network_error"


def command_fetch_verified(args: argparse.Namespace) -> int:
    args.manifest = configured_output_child(args, "manifest", "oa_candidates.csv")
    args.pdf_dir = configured_path(args, "pdf_dir", "pdf_dir")
    args.xml_dir = configured_path(args, "xml_dir", "xml_dir")
    args.log_dir = configured_output_child(args, "log_dir", "fetch")
    if args.delay < 3.0:
        print("--delay must be at least 3 seconds for full-text retrieval", file=sys.stderr)
        return 2
    rows = approved_rows(args.manifest)
    if args.limit:
        rows = rows[: args.limit]
    hosts = {host.casefold() for host in args.allow_host}
    plan = []
    for row in rows:
        parsed = urllib.parse.urlsplit(row.get("candidate_url", ""))
        plan.append({"record_id": row.get("record_id", ""), "host": parsed.hostname or "", "expected_filename": row.get("expected_filename", "")})
    if not args.execute:
        print(json.dumps({"mode": "dry_run", "approved_records": len(rows), "plan": plan, "required": "repeat with --execute after reviewing hosts"}, ensure_ascii=False, indent=2))
        return 0
    if not hosts:
        print("--execute requires at least one --allow-host", file=sys.stderr)
        return 2
    args.log_dir.mkdir(parents=True, exist_ok=True)
    event_path = args.log_dir / f"fetch_{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}.jsonl"
    status_counts: Counter[str] = Counter()
    consecutive_failures = 0
    stopped = False
    for row in rows:
        url = row.get("candidate_url", "")
        parsed = urllib.parse.urlsplit(url)
        host = (parsed.hostname or "").casefold()
        event: dict[str, Any] = {"timestamp_utc": utcnow(), "action": "fetch_verified", "record_id": row.get("record_id", ""), "host": host, "source_url_redacted": redact_url(url)}
        if parsed.scheme not in {"http", "https"} or not is_allowed_host(host, hosts):
            event["status"] = "host_not_allowed"
        else:
            suffix = ".xml" if row.get("expected_filename", "").casefold().endswith(".xml") else ".pdf"
            destination_dir = args.xml_dir if suffix == ".xml" else args.pdf_dir
            filename = row.get("expected_filename") or f"{row.get('record_id', 'download')}{suffix}"
            destination = destination_dir / Path(filename).name
            if destination.exists():
                event["status"] = "already_present"
            else:
                args.log_dir.mkdir(parents=True, exist_ok=True)
                temporary = args.log_dir / (destination.name + ".part")
                http_status, detail = download_one(url, temporary, args.timeout, args.max_bytes, hosts)
                event.update(http_status=http_status, transfer_detail=detail)
                if http_status in STOP_HTTP:
                    event["status"] = "rate_limited" if http_status == 429 else "challenge_or_block"
                    stopped = True
                elif http_status != 200 or not temporary.exists():
                    event["status"] = "network_error"
                else:
                    record = canonical_record(row, int(row.get("excel_row") or 2))
                    if suffix == ".pdf":
                        valid, validation, _ = validate_pdf(temporary, record, args.minimum_pdf_bytes)
                    else:
                        valid, validation = validate_xml(temporary, record, args.minimum_pdf_bytes)
                    event["validation"] = validation
                    if not valid:
                        event["status"] = "identity_mismatch"
                    else:
                        destination_dir.mkdir(parents=True, exist_ok=True)
                        temporary.replace(destination)
                        event.update(status="valid_xml" if suffix == ".xml" else "valid_pdf", bytes=destination.stat().st_size, sha256=sha256_file(destination), output_file=destination.name)
                if temporary.exists():
                    temporary.unlink()
        append_jsonl(event_path, event)
        status_counts[event["status"]] += 1
        if event["status"] in SUCCESS_STATUSES:
            consecutive_failures = 0
        else:
            consecutive_failures += 1
        if stopped or consecutive_failures >= args.max_consecutive_failures:
            stopped = True
            break
        time.sleep(max(0.0, args.delay))
    summary = {"finished_utc": utcnow(), "requested": len(rows), "status": dict(status_counts), "circuit_breaker_triggered": stopped, "events": event_path.name}
    write_json(args.log_dir / "summary.json", summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 3 if stopped else 0


def load_config(path: Path) -> dict[str, Any]:
    if tomllib is None:
        raise RuntimeError("Python 3.11 or newer is required")
    with path.open("rb") as handle:
        return tomllib.load(handle)


def command_doctor(args: argparse.Namespace) -> int:
    config = load_config(args.config)
    base = args.config.resolve().parent
    credential_names = list((config.get("credentials") or {}).values())
    project = config.get("project") or {}
    project_paths = {}
    for key in ("input", "pdf_dir", "xml_dir", "output_root", "manual_inbox"):
        value = project.get(key)
        path = resolve_local_path(str(value), base) if value else None
        project_paths[key] = {
            "configured": bool(value),
            "exists": bool(path and path.exists()),
            "kind": "file" if path and path.is_file() else ("directory" if path and path.is_dir() else "missing"),
            "writable": bool(path and path.exists() and os.access(path, os.W_OK)),
        }
    result = {
        "helper_version": VERSION,
        "python": sys.version.split()[0],
        "python_supported": sys.version_info >= (3, 11),
        "dependencies": {"openpyxl": openpyxl is not None, "pypdf": PdfReader is not None},
        "credentials": {name: bool(os.getenv(str(name), "")) for name in credential_names},
        "config_sections": sorted(config),
        "project_paths": project_paths,
        "environment_ready": sys.version_info >= (3, 11) and openpyxl is not None and PdfReader is not None,
        "project_ready": project_paths["input"]["exists"] and all(project_paths[key]["exists"] for key in ("pdf_dir", "xml_dir", "output_root", "manual_inbox")),
        "network_tested": False,
    }
    if args.clash_config:
        text = args.clash_config.read_text(encoding="utf-8", errors="ignore")
        mode_match = re.search(r"(?m)^mode:\s*([^\s#]+)", text)
        result["clash"] = {
            "mode": mode_match.group(1) if mode_match else "unknown",
            "openai_rule_present": bool(re.search(r"(?i)openai\.com|chatgpt\.com", text)),
            "institution_direct_rule_present": bool(re.search(r"(?i)(?:whu\.edu\.cn|metaersp\.com).*,\s*DIRECT", text)),
            "publisher_direct_rule_present": bool(re.search(r"(?i)(?:sciencedirect|tandfonline|ieee|wiley).*,\s*DIRECT", text)),
        }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    ok = result["environment_ready"]
    return 0 if ok else 2


def command_self_test(args: argparse.Namespace) -> int:
    checks: dict[str, bool] = {}
    with tempfile.TemporaryDirectory() as temp:
        root = Path(temp)
        bibliography = root / "bibliography.csv"
        with bibliography.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["Article Title", "DOI", "Publication Year", "Publisher"])
            writer.writeheader()
            writer.writerow({
                "Article Title": "Synthetic Fulltext Fetcher Self Test",
                "DOI": "10.5555/example.selftest", "Publication Year": "2025",
                "Publisher": "Example Publisher",
            })
        records = load_records(bibliography)
        checks["csv_input"] = len(records) == 1 and records[0]["doi"] == "10.5555/example.selftest"

        pdf_dir = root / "papers_pdf"
        xml_dir = root / "papers_xml"
        pdf_dir.mkdir()
        xml_dir.mkdir()
        pdf = pdf_dir / records[0]["expected_filename"]
        pdf.write_bytes(b"%PDF-1.4\n" + b"synthetic-image-payload\n" * 20)
        approval = root / "approval.jsonl"
        approval.write_text(json.dumps({
            "status": "visual_approved_imported", "expected_filename": pdf.name,
            "sha256": sha256_file(pdf),
        }) + "\n", encoding="utf-8")
        audited, anomalies, summary = audit_records(records, pdf_dir, xml_dir, 20, load_visual_approvals([approval]))
        checks["audit_and_visual_approval"] = summary["fulltext_coverage"] == 1 and not anomalies and audited[0]["fulltext_file"] == pdf.name

        missing = root / "missing.csv"
        write_csv(missing, [{**records[0], "status": "missing"}])
        manual_dir = root / "manual"
        with contextlib.redirect_stdout(io.StringIO()):
            manual_code = command_manual_queue(argparse.Namespace(remaining=missing, output_dir=manual_dir, config=None))
        checks["manual_queue"] = manual_code == 0 and (manual_dir / "manual_queue.csv").exists() and (manual_dir / "manual_queue.html").exists()

        clean_source = root / "clean-source"
        clean_source.mkdir()
        (clean_source / "README.md").write_text("Synthetic public fixture\n", encoding="utf-8")
        checks["privacy_scan"] = not privacy_findings(clean_source)

        if openpyxl is not None:
            workbook_path = root / "bibliography.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["Article Title", "DOI", "Publication Year", "Publisher"])
            sheet.append(["Synthetic Excel Self Test", "10.5555/example.xlsx", 2025, "Example Publisher"])
            workbook.save(workbook_path)
            workbook.close()
            checks["excel_input"] = load_records(workbook_path)[0]["doi"] == "10.5555/example.xlsx"
        else:
            checks["excel_input"] = False

        if PdfReader is not None:
            from pypdf import PdfWriter

            blank = root / "blank.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=300, height=300)
            with blank.open("wb") as handle:
                writer.write(handle)
            blank_record = {**records[0], "expected_filename": blank.name}
            valid, detail, _ = validate_pdf(blank, blank_record, 100)
            checks["pdf_identity_validation"] = not valid and "identity_mismatch" in detail
        else:
            checks["pdf_identity_validation"] = False

    result = {
        "finished_utc": utcnow(), "offline": True, "synthetic_data_only": True,
        "checks": checks, "passed": all(checks.values()),
    }
    if args.output:
        write_json(args.output, result)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["passed"] else 2


def privacy_findings(root: Path) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    content_patterns = [
        ("windows_user_path", re.compile(r"(?i)[A-Z]:[\\/]Users[\\/](?!Public|Default|<user>|<name>)[^\\/\s'\"`]+")),
        ("posix_user_path", re.compile(r"/(?:home|Users)/(?!<user>|<name>)[^/\s'\"`]+")),
        ("personal_email", re.compile(r"(?i)\b[A-Z0-9._%+-]+@(?!example\.(?:com|org|edu)\b)[A-Z0-9.-]+\.[A-Z]{2,}\b")),
        ("env_secret_assignment", re.compile(r"(?i)(?:\$env:|setx\s+)[A-Z0-9_]*(?:KEY|TOKEN|SECRET|PASSWORD)[A-Z0-9_]*\s*(?:=|\s)\s*['\"](?!<|\$)[^'\"]{8,}")),
        ("authorization_value", re.compile(r"(?i)\b(?:Authorization|Cookie)\s*[:=]\s*['\"](?!<)[^'\"]{8,}")),
        ("signed_url", re.compile(r"(?i)https?://[^\s'\"]+[?&](?:token|signature|sig|key|auth|expires)=[^\s'\"&]+")),
    ]
    for path in root.rglob("*"):
        try:
            relative = path.relative_to(root)
        except ValueError:
            continue
        if any(part in SKIP_DIRS for part in relative.parts):
            continue
        if path.is_dir():
            if path.name.casefold() in {"browser_profiles", "user-data-dir", "papers_pdf", "papers_xml", "outputs", "manual_pdf_inbox"}:
                findings.append({"path": relative.as_posix(), "line": 0, "kind": "private_or_copyright_directory"})
            continue
        suffix = path.suffix.casefold()
        if suffix in RISKY_SUFFIXES:
            findings.append({"path": relative.as_posix(), "line": 0, "kind": "risky_file_type"})
            continue
        if suffix == ".csv" and re.search(r"(?i)(bibliography|savedrecs|corpus[_-]?audit|download[_-]?log|manifest)", path.name):
            findings.append({"path": relative.as_posix(), "line": 0, "kind": "possible_research_or_run_data"})
            continue
        if suffix not in TEXT_SUFFIXES and path.name not in {"SKILL.md", "README.md", ".gitignore"}:
            continue
        try:
            if path.stat().st_size > 2_000_000:
                continue
            lines = path.read_text(encoding="utf-8").splitlines()
        except (OSError, UnicodeDecodeError):
            continue
        for number, line in enumerate(lines, 1):
            for kind, pattern in content_patterns:
                if pattern.search(line):
                    findings.append({"path": relative.as_posix(), "line": number, "kind": kind})
    return findings


def command_privacy_scan(args: argparse.Namespace) -> int:
    findings = privacy_findings(args.path.resolve())
    result = {"scanned_at_utc": utcnow(), "root": args.path.resolve().name, "findings": findings, "finding_count": len(findings)}
    if args.output:
        write_json(args.output, result)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if not findings else 3


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Quick start:\n"
            "  paper_fulltext_fetcher.py init --workspace my-review\n"
            "  paper_fulltext_fetcher.py self-test\n"
            "  paper_fulltext_fetcher.py doctor --config my-review/acquisition.local.toml\n"
            "  paper_fulltext_fetcher.py inspect-input --config my-review/acquisition.local.toml\n"
            "  paper_fulltext_fetcher.py audit --config my-review/acquisition.local.toml"
        ),
    )
    parser.add_argument("--version", action="version", version=VERSION)
    sub = parser.add_subparsers(dest="command", required=True)

    initializer = sub.add_parser("init", help="Create a privacy-safe project workspace and local config")
    initializer.add_argument("--workspace", type=Path, default=Path("."))
    initializer.add_argument("--input", type=Path, default=Path("bibliography.xlsx"), help="Relative bibliography path inside the workspace")
    initializer.add_argument("--dry-run", action="store_true")
    initializer.add_argument("--force", action="store_true", help="Overwrite acquisition.local.toml only; never deletes corpus files")
    initializer.set_defaults(function=command_init)

    doctor = sub.add_parser("doctor", help="Check local dependencies and configured credential names without printing values")
    doctor.add_argument("--config", type=Path, default=skill_root() / "assets" / "acquisition.example.toml")
    doctor.add_argument("--clash-config", type=Path)
    doctor.set_defaults(function=command_doctor)

    self_test = sub.add_parser("self-test", help="Run an offline end-to-end check using synthetic data only")
    self_test.add_argument("--output", type=Path)
    self_test.set_defaults(function=command_self_test)

    inspector = sub.add_parser("inspect-input", help="Check spreadsheet/sheet/column compatibility without printing bibliographic content")
    inspector.add_argument("--config", type=Path)
    inspector.add_argument("--input", type=Path)
    inspector.add_argument("--sheet")
    inspector.add_argument("--column", action="append", default=[], metavar="FIELD=HEADER")
    inspector.add_argument("--output", type=Path)
    inspector.set_defaults(function=command_inspect_input)

    audit = sub.add_parser("audit", help="Audit bibliography coverage using local PDF/XML files")
    audit.add_argument("--config", type=Path)
    audit.add_argument("--input", type=Path)
    audit.add_argument("--sheet")
    audit.add_argument("--column", action="append", default=[], metavar="FIELD=HEADER")
    audit.add_argument("--pdf-dir", type=Path)
    audit.add_argument("--xml-dir", type=Path)
    audit.add_argument("--output-dir", type=Path)
    audit.add_argument("--minimum-pdf-bytes", type=int, default=10_000)
    audit.add_argument("--visual-approvals", type=Path, action="append", default=[])
    audit.set_defaults(function=command_audit)

    manual = sub.add_parser("manual-queue", help="Build CSV and HTML queues for lawful manual acquisition")
    manual.add_argument("--config", type=Path)
    manual.add_argument("--remaining", type=Path)
    manual.add_argument("--output-dir", type=Path)
    manual.set_defaults(function=command_manual_queue)

    importer = sub.add_parser("import", help="Validate and copy user-supplied PDFs without deleting originals")
    importer.add_argument("--config", type=Path)
    importer.add_argument("--input", type=Path)
    importer.add_argument("--sheet")
    importer.add_argument("--column", action="append", default=[], metavar="FIELD=HEADER")
    importer.add_argument("--remaining", type=Path)
    importer.add_argument("--inbox", type=Path)
    importer.add_argument("--pdf-dir", type=Path)
    importer.add_argument("--log-dir", type=Path)
    importer.add_argument("--minimum-pdf-bytes", type=int, default=10_000)
    importer.add_argument("--dry-run", action="store_true")
    importer.set_defaults(function=command_import)

    discovery = sub.add_parser("discover-oa", help="Discover OA candidates; never downloads full text")
    discovery.add_argument("--config", type=Path)
    discovery.add_argument("--remaining", type=Path)
    discovery.add_argument("--output", type=Path)
    discovery.add_argument("--unpaywall-email-env", default="UNPAYWALL_EMAIL")
    discovery.add_argument("--openalex-key-env", default="OPENALEX_API_KEY")
    discovery.add_argument("--delay", type=float, default=5.0)
    discovery.add_argument("--timeout", type=int, default=30)
    discovery.add_argument("--limit", type=int, default=1)
    discovery.set_defaults(function=command_discover_oa)

    fetch = sub.add_parser("fetch-verified", help="Fetch only explicitly approved OA/API URLs; dry-run unless --execute")
    fetch.add_argument("--config", type=Path)
    fetch.add_argument("--manifest", type=Path)
    fetch.add_argument("--pdf-dir", type=Path)
    fetch.add_argument("--xml-dir", type=Path)
    fetch.add_argument("--log-dir", type=Path)
    fetch.add_argument("--allow-host", action="append", default=[])
    fetch.add_argument("--delay", type=float, default=10.0)
    fetch.add_argument("--timeout", type=int, default=60)
    fetch.add_argument("--max-bytes", type=int, default=100 * 1024 * 1024)
    fetch.add_argument("--minimum-pdf-bytes", type=int, default=10_000)
    fetch.add_argument("--max-consecutive-failures", type=int, default=2)
    fetch.add_argument("--limit", type=int, default=1)
    fetch.add_argument("--execute", action="store_true")
    fetch.set_defaults(function=command_fetch_verified)

    privacy = sub.add_parser("privacy-scan", help="Report risky files and secret-like content without printing matched values")
    privacy.add_argument("--path", type=Path, required=True)
    privacy.add_argument("--output", type=Path)
    privacy.set_defaults(function=command_privacy_scan)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    try:
        return int(args.function(args))
    except KeyboardInterrupt:
        print("Interrupted; existing event logs are preserved.", file=sys.stderr)
        return 130
    except Exception as exc:
        print(f"{type(exc).__name__}: {redact_error_message(str(exc))}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
